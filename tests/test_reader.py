from __future__ import annotations

from io import BytesIO

from docx import Document
from reportlab.pdfgen import canvas

from org_audit.reader import read_uploaded_file


def make_pdf(text: str | None = "Network team checks reserve equipment") -> bytes:
    output = BytesIO()
    page = canvas.Canvas(output)
    if text:
        page.drawString(72, 720, text)
    page.save()
    return output.getvalue()


def make_docx() -> bytes:
    document = Document()
    document.add_heading("Отдел эксплуатации сети", level=1)
    document.add_paragraph("1. Проверяет резервное оборудование.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Ответственный"
    table.cell(0, 1).text = "Главный инженер"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_txt_keeps_locations_units_and_stable_ids():
    content = "Отдел эксплуатации сети\n1. Проверяет узлы связи.\n\n2. Фиксирует неисправности.\n"
    first = read_uploaded_file("source.txt", content.encode("utf-8"), "ДО", 0)
    second = read_uploaded_file("source.txt", content.encode("utf-8"), "ДО", 0)

    assert first.status == "ok"
    assert len(first.fragments) == 3
    assert first.fragments[0].text == "Отдел эксплуатации сети"
    assert first.fragments[1].location == "текст, строка 2"
    assert first.fragments[1].unit == "Отдел эксплуатации сети"
    assert [item.id for item in first.fragments] == [item.id for item in second.fragments]


def test_docx_reads_paragraph_number_and_table_location():
    result = read_uploaded_file("policy.docx", make_docx(), "ПОСЛЕ")

    assert result.status == "ok"
    assert any(item.text == "1. Проверяет резервное оборудование." and item.location == "пункт 1" for item in result.fragments)
    table_fragment = next(item for item in result.fragments if "Ответственный" in item.text)
    assert table_fragment.location == "таблица 1, строка 1"
    assert table_fragment.unit == "Отдел эксплуатации сети"


def test_pdf_keeps_one_based_page_number():
    result = read_uploaded_file("policy.pdf", make_pdf(), "ДО")

    assert result.status == "ok"
    assert result.fragments
    assert result.fragments[0].location.startswith("страница 1")
    assert "Network team" in result.fragments[0].text


def test_mixed_pdf_pages_are_marked_partial():
    output = BytesIO()
    page = canvas.Canvas(output)
    page.drawString(72, 720, "Page with searchable text")
    page.showPage()
    page.showPage()
    page.save()

    result = read_uploaded_file("mixed.pdf", output.getvalue(), "ДО")

    assert result.status == "partial"
    assert result.fragments
    assert "страницах 2" in result.message
    assert "OCR" in result.message


def test_scan_empty_and_damaged_files_are_reported():
    scanned = read_uploaded_file("scan.pdf", make_pdf(None), "ДО")
    blank_txt = read_uploaded_file("blank.txt", b" \n ", "ДО")
    damaged_docx = read_uploaded_file("broken.docx", b"not a docx", "ДО")
    damaged_pdf = read_uploaded_file("broken.pdf", b"not a pdf", "ДО")

    assert scanned.status == "error" and "OCR" in scanned.message
    assert blank_txt.status == "error" and "нет текста" in blank_txt.message
    assert damaged_docx.status == "error"
    assert damaged_pdf.status == "error"


def test_empty_and_unsupported_files_are_clear_errors():
    empty = read_uploaded_file("empty.docx", b"", "ПОСЛЕ")
    unsupported = read_uploaded_file("data.csv", b"x,y", "ПОСЛЕ")

    assert empty.status == "error" and "пуст" in empty.message
    assert unsupported.status == "error" and "не поддерживается" in unsupported.message

